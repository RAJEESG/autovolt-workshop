import os
import sys
import json
import uuid
import io
import re
from datetime import datetime, date, timedelta
from typing import Optional, List, Dict, Any

from fastapi import FastAPI, Request, Form, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database as db
from auth_helper import hash_password, verify_password, create_signed_session_cookie, verify_signed_session_cookie

load_dotenv()

# Initialize DB
db.init_db()

app = FastAPI(
    title="AutoVolt Cloud HQ — SaaS Master Vendor Platform",
    description="Dedicated Standalone Master Management Application for AutoVolt Pro SaaS Platform",
    version="2.0.0"
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
if os.path.exists(STATIC_DIR):
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))

def format_currency_filter(val):
    try:
        if val is None or str(val).strip() == "":
            return "0.00"
        return f"{float(val):,.2f}"
    except Exception:
        return "0.00"

templates.env.filters["currency"] = format_currency_filter

# =========================================================
# 🔐 AUTHENTICATION MIDDLEWARE FOR SAAS MASTER VENDOR PORTAL
# =========================================================
@app.middleware("http")
async def vendor_auth_middleware(request: Request, call_next):
    path = request.url.path
    exempt_prefixes = ["/static", "/login", "/logout"]
    is_exempt = any(path.startswith(p) for p in exempt_prefixes)
    
    raw_cookie = request.cookies.get("vendor_session")
    session_user = verify_signed_session_cookie(raw_cookie)
    if not is_exempt and not session_user and request.method == "GET":
        return RedirectResponse(url="/login", status_code=303)
        
    if path == "/login" and session_user and request.method == "GET":
        return RedirectResponse(url="/", status_code=303)
        
    response = await call_next(request)
    return response

# =========================================================
# 🔑 LOGIN & LOGOUT
# =========================================================
@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request, error: Optional[str] = None):
    return templates.TemplateResponse(
        request=request,
        name="admin_hq/login.html",
        context={"error": error}
    )

@app.post("/login")
async def handle_login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...)
):
    conn = db.get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM users WHERE username = ? AND status = 'ACTIVE'", (username.strip(),))
    user = cursor.fetchone()

    if not user or not verify_password(password.strip(), user["password_hash"]):
        conn.close()
        return templates.TemplateResponse(
            request=request,
            name="admin_hq/login.html",
            context={"error": "Invalid Master Admin credentials. Please try again."}
        )

    # Allow superadmin or admin
    if user["role"] not in ("SUPER_ADMIN", "ADMIN"):
        conn.close()
        return templates.TemplateResponse(
            request=request,
            name="admin_hq/login.html",
            context={"error": "Access Denied: Only SaaS Master Vendor Administrators can enter."}
        )

    db.log_audit("VENDOR_HQ_LOGIN", "users", user["id"], f"Master admin {user['username']} logged into HQ", cursor=cursor)
    conn.commit()
    conn.close()

    response = RedirectResponse(url="/", status_code=303)
    signed_val = create_signed_session_cookie(user["username"], expires_in=86400 * 30)
    response.set_cookie(key="vendor_session", value=signed_val, max_age=86400 * 30, httponly=True, samesite="lax")
    return response

@app.get("/logout")
async def handle_logout():
    response = RedirectResponse(url="/login", status_code=303)
    response.delete_cookie("vendor_session")
    return response

# =========================================================
# 📊 SAAS MASTER DASHBOARD OVERVIEW
# =========================================================
@app.get("/", response_class=HTMLResponse)
async def dashboard_page(request: Request):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM tenants ORDER BY created_at DESC")
    tenants = [dict(row) for row in cursor.fetchall()]

    if not tenants:
        try:
            import supabase_sync
            supabase_sync.sync_critical_from_supabase()
            cursor.execute("SELECT * FROM tenants ORDER BY created_at DESC")
            tenants = [dict(row) for row in cursor.fetchall()]
        except Exception as e:
            print(f"[Supabase Sync on Clients Page] {e}")

    today = datetime.now().date()
    today_str = today.strftime("%Y-%m-%d")
    thirty_days_later = (today + timedelta(days=30)).strftime("%Y-%m-%d")

    active_tenants = 0
    expired_tenants = 0
    renewals_due_count = 0

    for t in tenants:
        exp_date_str = t.get("subscription_end_date", "")
        if exp_date_str:
            try:
                exp_date = datetime.strptime(exp_date_str, "%Y-%m-%d").date()
                days_left = (exp_date - today).days
                t["is_active"] = (t.get("status") == "ACTIVE") and (days_left >= 0)
                if days_left < 0:
                    t["days_left_text"] = f"Expired {abs(days_left)} days ago"
                elif days_left == 0:
                    t["days_left_text"] = "Expires today!"
                else:
                    t["days_left_text"] = f"{days_left} days remaining"

                if 0 <= days_left <= 30:
                    renewals_due_count += 1
            except Exception:
                t["is_active"] = (t.get("status") == "ACTIVE")
                t["days_left_text"] = "Active"
        else:
            t["is_active"] = (t.get("status") == "ACTIVE")
            t["days_left_text"] = "No expiry date"

        if t["is_active"]:
            active_tenants += 1
        else:
            expired_tenants += 1

    # Financial Stats from SaaS Vendor Invoices
    cursor.execute("SELECT COALESCE(SUM(total_amount), 0), COUNT(*) FROM saas_vendor_invoices WHERE payment_status = 'PAID'")
    rev_row = cursor.fetchone()
    total_revenue = rev_row[0]

    # Calculate MRR estimate
    cursor.execute("SELECT plan_name, billing_cycle, price_paid FROM subscriptions WHERE is_active = 1")
    subs = cursor.fetchall()
    mrr = 0.0
    for s in subs:
        price = s["price_paid"] or 0.0
        if s["billing_cycle"] == "YEARLY":
            mrr += (price / 12.0)
        else:
            mrr += price

    cursor.execute("SELECT * FROM saas_vendor_invoices ORDER BY created_at DESC LIMIT 6")
    recent_invoices = [dict(row) for row in cursor.fetchall()]

    conn.close()

    stats = {
        "total_tenants": len(tenants),
        "active_tenants": active_tenants,
        "expired_tenants": expired_tenants,
        "renewals_due_count": renewals_due_count,
        "total_revenue": total_revenue,
        "mrr": mrr
    }

    return templates.TemplateResponse(
        request=request,
        name="admin_hq/dashboard.html",
        context={
            "active_page": "dashboard",
            "stats": stats,
            "tenants": tenants[:6],
            "recent_invoices": recent_invoices,
            "total_tenants": len(tenants)
        }
    )

# =========================================================
# 🏢 CLIENT WORKSHOPS & MULTI-TENANT CONFIGURATION
# =========================================================
@app.get("/clients", response_class=HTMLResponse)
async def clients_page(request: Request, msg: Optional[str] = None):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM tenants ORDER BY created_at DESC")
    tenants = [dict(row) for row in cursor.fetchall()]

    if not tenants:
        try:
            import supabase_sync
            supabase_sync.sync_critical_from_supabase()
            cursor.execute("SELECT * FROM tenants ORDER BY created_at DESC")
            tenants = [dict(row) for row in cursor.fetchall()]
        except Exception as e:
            print(f"[Supabase Sync on Clients Page] {e}")

    today = datetime.now().date()
    for t in tenants:
        exp_str = t.get("subscription_end_date", "")
        if exp_str:
            try:
                exp_date = datetime.strptime(exp_str, "%Y-%m-%d").date()
                t["is_active"] = (t.get("status") == "ACTIVE") and (exp_date >= today)
            except Exception:
                t["is_active"] = (t.get("status") == "ACTIVE")
        else:
            t["is_active"] = (t.get("status") == "ACTIVE")

        # Parse GST Slabs JSON
        try:
            t["gst_slabs_list"] = json.loads(t.get("gst_slabs") or '["0", "5", "12", "18", "28"]')
        except Exception:
            t["gst_slabs_list"] = ["0", "5", "12", "18", "28"]

    conn.close()

    return templates.TemplateResponse(
        request=request,
        name="admin_hq/clients.html",
        context={
            "active_page": "clients",
            "tenants": tenants,
            "msg": msg,
            "total_tenants": len(tenants)
        }
    )

@app.post("/api/admin/clients/create")
async def create_client_workshop(
    request: Request,
    name: str = Form(...),
    owner_name: str = Form(...),
    whatsapp_mobile: str = Form(...),
    email: Optional[str] = Form(None),
    address: str = Form(...),
    gstin: Optional[str] = Form(None),
    logo_url: Optional[str] = Form(None),
    client_code: Optional[str] = Form(None),
    upi_id: str = Form(...),
    upi_payee_name: Optional[str] = Form(None),
    bank_name: Optional[str] = Form(None),
    bank_account_no: Optional[str] = Form(None),
    bank_ifsc: Optional[str] = Form(None),
    language: str = Form("en"),
    master_pin: str = Form("1234"),
    initial_username: str = Form(...),
    initial_password: str = Form(...),
    plan_type: str = Form("YEARLY")
):
    form_data = await request.form()
    slabs_raw = form_data.getlist("gst_slabs")
    slabs_list = [str(s).strip() for s in slabs_raw if str(s).strip()]
    if not slabs_list:
        slabs_list = ["0", "5", "12", "18", "28"]
    gst_slabs_json = json.dumps(slabs_list)

    conn = db.get_db_connection()
    cursor = conn.cursor()

    tenant_id = f"tenant_{uuid.uuid4().hex[:8]}"
    start_date = datetime.now().strftime("%Y-%m-%d")
    days = 365 if plan_type == "YEARLY" else 30
    expiry_date = (datetime.now() + timedelta(days=days)).strftime("%Y-%m-%d")

    base_price = 9999.0 if plan_type == "YEARLY" else 999.0
    cgst = base_price * 0.09
    sgst = base_price * 0.09
    total_price = base_price + cgst + sgst

    clean_code = (client_code or "").strip().upper()
    if not clean_code:
        clean_code = re.sub(r'[^A-Z0-9]', '', name.upper())[:8] or f"WS{uuid.uuid4().hex[:4].upper()}"

    # Insert Tenant
    cursor.execute("""
    INSERT INTO tenants (
        id, client_code, name, owner_name, phone, whatsapp_mobile, email, address, gstin, logo_url,
        upi_id, upi_payee_name, bank_name, bank_account_no, bank_ifsc, bank_branch,
        gst_slabs, language, master_pin, status, subscription_plan, subscription_start_date, subscription_end_date, price_paid, created_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Main', ?, ?, ?, 'ACTIVE', ?, ?, ?, ?, datetime('now'))
    """, (
        tenant_id, clean_code, name.strip(), owner_name.strip(), whatsapp_mobile.strip(), whatsapp_mobile.strip(),
        email.strip() if email else "", address.strip(), gstin.strip() if gstin else "", logo_url.strip() if logo_url else "",
        upi_id.strip(), upi_payee_name.strip() if upi_payee_name else name.strip(),
        bank_name.strip() if bank_name else "", bank_account_no.strip() if bank_account_no else "",
        bank_ifsc.strip() if bank_ifsc else "", gst_slabs_json, language, master_pin.strip(),
        plan_type, start_date, expiry_date, total_price
    ))

    # Insert Subscription Record
    sub_id = f"sub_{uuid.uuid4().hex[:8]}"
    cursor.execute("""
    INSERT INTO subscriptions (id, tenant_id, plan_name, start_date, end_date, is_active, billing_cycle, price_paid, notes, created_at)
    VALUES (?, ?, ?, ?, ?, 1, ?, ?, 'Initial Onboarding License', datetime('now'))
    """, (sub_id, tenant_id, f"PRO_{plan_type}", start_date, expiry_date, plan_type, total_price))

    # Insert Workshop Profile Record for new Tenant
    cursor.execute("""
    INSERT INTO workshop_profile (
        id, tenant_id, client_code, name, subtitle, owner_name, phone, email, address, gstin, logo_url,
        upi_id, upi_payee_name, bank_name, bank_account_no, bank_ifsc, bank_branch,
        gst_slabs, language, currency, master_pin, terms, default_state, default_gst_type, tax_regime, round_off_enabled, created_at, updated_at
    ) VALUES (
        ?, ?, ?, ?, 'Automotive Electrical & Electronic Service Center', ?, ?, ?, ?, ?, ?,
        ?, ?, ?, ?, ?, 'Main',
        ?, ?, '₹', ?, '1. Electrical parts warranty as per manufacturer policy.\n2. Goods once fitted cannot be returned.\n3. Payment due upon completion.', 'Kerala (32)', 'INTRA_STATE', 'REGULAR_GST', 1, datetime('now'), datetime('now')
    )
    """, (
        tenant_id, tenant_id, clean_code, name.strip(), owner_name.strip(), whatsapp_mobile.strip(), email.strip() if email else "",
        address.strip(), gstin.strip() if gstin else "", logo_url.strip() if logo_url else "",
        upi_id.strip(), upi_payee_name.strip() if upi_payee_name else name.strip(),
        bank_name.strip() if bank_name else "", bank_account_no.strip() if bank_account_no else "",
        bank_ifsc.strip() if bank_ifsc else "", gst_slabs_json, language, master_pin.strip()
    ))

    # Create Initial Admin User for this Garage (isolated to this tenant_id)
    cursor.execute("SELECT id FROM users WHERE tenant_id = ? AND username = ?", (tenant_id, initial_username.strip()))
    existing_user = cursor.fetchone()
    pw_hash = hash_password(initial_password.strip())
    if not existing_user:
        user_id = f"user_{uuid.uuid4().hex[:8]}"
        cursor.execute("""
        INSERT INTO users (id, username, password_hash, full_name, role, email, phone, whatsapp_mobile, tenant_id, status, created_at)
        VALUES (?, ?, ?, ?, 'ADMIN', ?, ?, ?, ?, 'ACTIVE', datetime('now'))
        """, (
            user_id, initial_username.strip(), pw_hash, owner_name.strip(),
            email.strip() if email else "", whatsapp_mobile.strip(), whatsapp_mobile.strip(), tenant_id
        ))
    else:
        cursor.execute("UPDATE users SET password_hash = ? WHERE id = ?", (pw_hash, existing_user["id"]))

    # Auto-generate SaaS Vendor Invoice for the user's company books
    inv_id = f"sinv_{uuid.uuid4().hex[:8]}"
    cursor.execute("SELECT COUNT(*) FROM saas_vendor_invoices")
    inv_count = cursor.fetchone()[0] + 1
    inv_no = f"HQ-INV-2026-{inv_count:03d}"
    cursor.execute("SELECT id FROM saas_vendor_invoices WHERE invoice_no = ?", (inv_no,))
    while cursor.fetchone():
        inv_count += 1
        inv_no = f"HQ-INV-2026-{inv_count:03d}"
        cursor.execute("SELECT id FROM saas_vendor_invoices WHERE invoice_no = ?", (inv_no,))

    cursor.execute("""
    INSERT INTO saas_vendor_invoices (
        id, invoice_no, tenant_id, tenant_name, invoice_date, plan_type,
        subscription_start, subscription_end, taxable_amount, cgst_rate, cgst_amount,
        sgst_rate, sgst_amount, total_amount, payment_status, payment_mode, notes, created_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 9.0, ?, 9.0, ?, ?, 'PAID', 'UPI', 'Initial Workshop Setup License Invoice', datetime('now'))
    """, (
        inv_id, inv_no, tenant_id, name.strip(), start_date, plan_type,
        start_date, expiry_date, base_price, cgst, sgst, total_price
    ))

    # Auto-seed Default Cash Account & Bank Account for new Workshop
    acc_cash_id = f"acc_{uuid.uuid4().hex[:8]}"
    cursor.execute("""
    INSERT INTO bank_accounts (id, tenant_id, account_name, account_type, current_balance, is_primary, created_at)
    VALUES (?, ?, 'Daily Cash Register', 'CASH', 0.0, 1, datetime('now'))
    """, (acc_cash_id, tenant_id))

    acc_bank_id = None
    if bank_name:
        acc_bank_id = f"acc_{uuid.uuid4().hex[:8]}"
        cursor.execute("""
        INSERT INTO bank_accounts (id, tenant_id, account_name, bank_name, account_number, ifsc_code, upi_vpa, account_type, current_balance, is_primary, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, 'BANK', 0.0, 0, datetime('now'))
        """, (acc_bank_id, tenant_id, f"{bank_name.strip()} Operating Account", bank_name.strip(), bank_account_no or "", bank_ifsc or "", upi_id.strip()))

    # Auto-seed Lead Technician / Staff Member
    tech_id = f"tech_{uuid.uuid4().hex[:8]}"
    cursor.execute("""
    INSERT INTO technicians (id, tenant_id, name, phone, role, status, created_at)
    VALUES (?, ?, ?, ?, 'Lead Electrician', 'Active', datetime('now'))
    """, (tech_id, tenant_id, owner_name.strip(), whatsapp_mobile.strip()))

    db.log_audit("ONBOARD_CLIENT_WORKSHOP", "tenants", tenant_id, f"Onboarded client garage {name} with {plan_type} plan", cursor=cursor)
    conn.commit()
    conn.close()

    # Real-Time Cloud Sync: Push newly created workshop & admin credentials to Supabase
    try:
        import supabase_sync
        supabase_sync.push_to_supabase("tenants", {
            "id": tenant_id, "client_code": clean_code, "name": name.strip(),
            "owner_name": owner_name.strip(), "phone": whatsapp_mobile.strip(),
            "whatsapp_mobile": whatsapp_mobile.strip(), "email": email.strip() if email else "",
            "address": address.strip(), "gstin": gstin.strip() if gstin else "",
            "logo_url": logo_url.strip() if logo_url else "", "upi_id": upi_id.strip(),
            "upi_payee_name": upi_payee_name.strip() if upi_payee_name else name.strip(),
            "bank_name": bank_name.strip() if bank_name else "",
            "bank_account_no": bank_account_no.strip() if bank_account_no else "",
            "bank_ifsc": bank_ifsc.strip() if bank_ifsc else "", "bank_branch": "Main",
            "gst_slabs": gst_slabs_json, "language": language, "master_pin": master_pin.strip(),
            "status": "ACTIVE", "subscription_plan": plan_type, "subscription_start_date": start_date,
            "subscription_end_date": expiry_date, "price_paid": total_price
        })
        supabase_sync.push_to_supabase("subscriptions", {
            "id": sub_id, "tenant_id": tenant_id, "plan_name": f"PRO_{plan_type}",
            "start_date": start_date, "end_date": expiry_date, "is_active": 1,
            "billing_cycle": plan_type, "price_paid": total_price, "notes": "Initial Onboarding License"
        })
        supabase_sync.push_to_supabase("workshop_profile", {
            "id": tenant_id, "tenant_id": tenant_id, "client_code": clean_code, "name": name.strip(),
            "subtitle": "Automotive Electrical & Electronic Service Center",
            "owner_name": owner_name.strip(), "phone": whatsapp_mobile.strip(),
            "email": email.strip() if email else "", "address": address.strip(),
            "gstin": gstin.strip() if gstin else "", "logo_url": logo_url.strip() if logo_url else "",
            "upi_id": upi_id.strip(), "upi_payee_name": upi_payee_name.strip() if upi_payee_name else name.strip(),
            "bank_name": bank_name.strip() if bank_name else "",
            "bank_account_no": bank_account_no.strip() if bank_account_no else "",
            "bank_ifsc": bank_ifsc.strip() if bank_ifsc else "", "bank_branch": "Main",
            "gst_slabs": gst_slabs_json, "language": language, "master_pin": master_pin.strip()
        })
        supabase_sync.push_to_supabase("users", {
            "id": user_id, "username": initial_username.strip(), "password_hash": pw_hash,
            "full_name": owner_name.strip(), "role": "ADMIN",
            "email": email.strip() if email else "", "phone": whatsapp_mobile.strip(),
            "whatsapp_mobile": whatsapp_mobile.strip(), "tenant_id": tenant_id, "status": "ACTIVE"
        })
        supabase_sync.push_to_supabase("saas_vendor_invoices", {
            "id": inv_id, "invoice_no": inv_no, "tenant_id": tenant_id, "tenant_name": name.strip(),
            "invoice_date": start_date, "plan_type": plan_type, "subscription_start": start_date,
            "subscription_end": expiry_date, "taxable_amount": base_price, "cgst_rate": 9.0,
            "cgst_amount": cgst, "sgst_rate": 9.0, "sgst_amount": sgst, "total_amount": total_price,
            "payment_status": "PAID", "payment_mode": "UPI", "notes": "Initial Workshop Setup License Invoice"
        })
    except Exception as e:
        print(f"[Supabase Onboarding Sync] {e}")

    return RedirectResponse(url="/clients?msg=Workshop+successfully+onboarded+and+license+activated!", status_code=303)

@app.post("/api/admin/clients/update")
async def update_client_workshop(
    request: Request,
    tenant_id: str = Form(...),
    name: str = Form(...),
    owner_name: str = Form(...),
    whatsapp_mobile: str = Form(...),
    email: Optional[str] = Form(None),
    address: str = Form(...),
    gstin: Optional[str] = Form(None),
    logo_url: Optional[str] = Form(None),
    client_code: Optional[str] = Form(None),
    upi_id: str = Form(...),
    upi_payee_name: Optional[str] = Form(None),
    bank_name: Optional[str] = Form(None),
    bank_account_no: Optional[str] = Form(None),
    bank_ifsc: Optional[str] = Form(None),
    language: str = Form("en"),
    master_pin: str = Form("1234")
):
    form_data = await request.form()
    slabs_raw = form_data.getlist("gst_slabs")
    slabs_list = [str(s).strip() for s in slabs_raw if str(s).strip()]
    if not slabs_list:
        slabs_list = ["0", "5", "12", "18", "28"]
    gst_slabs_json = json.dumps(slabs_list)

    conn = db.get_db_connection()
    cursor = conn.cursor()

    clean_code = (client_code or "").strip().upper()

    cursor.execute("""
    UPDATE tenants SET
        client_code = COALESCE(NULLIF(?, ''), client_code),
        name = ?, owner_name = ?, phone = ?, whatsapp_mobile = ?, email = ?,
        address = ?, gstin = ?, logo_url = ?, upi_id = ?, upi_payee_name = ?,
        bank_name = ?, bank_account_no = ?, bank_ifsc = ?,
        gst_slabs = ?, language = ?, master_pin = ?
    WHERE id = ?
    """, (
        clean_code,
        name.strip(), owner_name.strip(), whatsapp_mobile.strip(), whatsapp_mobile.strip(), email.strip() if email else "",
        address.strip(), gstin.strip() if gstin else "", logo_url.strip() if logo_url else "",
        upi_id.strip(), upi_payee_name.strip() if upi_payee_name else name.strip(),
        bank_name.strip() if bank_name else "", bank_account_no.strip() if bank_account_no else "",
        bank_ifsc.strip() if bank_ifsc else "", gst_slabs_json, language, master_pin.strip(), tenant_id
    ))

    # Also update workshop_profile for this tenant
    cursor.execute("""
    UPDATE workshop_profile SET
        client_code = COALESCE(NULLIF(?, ''), client_code),
        name = ?, owner_name = ?, phone = ?, email = ?, address = ?, gstin = ?,
        logo_url = ?, upi_id = ?, upi_payee_name = ?, bank_name = ?, bank_account_no = ?, bank_ifsc = ?,
        gst_slabs = ?, language = ?, master_pin = ?, updated_at = datetime('now')
    WHERE tenant_id = ? OR id = ?
    """, (
        clean_code,
        name.strip(), owner_name.strip(), whatsapp_mobile.strip(), email.strip() if email else "",
        address.strip(), gstin.strip() if gstin else "", logo_url.strip() if logo_url else "",
        upi_id.strip(), upi_payee_name.strip() if upi_payee_name else name.strip(),
        bank_name.strip() if bank_name else "", bank_account_no.strip() if bank_account_no else "",
        bank_ifsc.strip() if bank_ifsc else "", gst_slabs_json, language, master_pin.strip(),
        tenant_id, tenant_id
    ))

    db.log_audit("UPDATE_CLIENT_WORKSHOP", "tenants", tenant_id, f"Updated configuration for {name}", cursor=cursor)
    conn.commit()
    conn.close()

    try:
        import supabase_sync
        supabase_sync.push_to_supabase("tenants", {
            "id": tenant_id, "client_code": clean_code, "name": name.strip(),
            "owner_name": owner_name.strip(), "phone": whatsapp_mobile.strip(),
            "whatsapp_mobile": whatsapp_mobile.strip(), "email": email.strip() if email else "",
            "address": address.strip(), "gstin": gstin.strip() if gstin else "",
            "logo_url": logo_url.strip() if logo_url else "", "upi_id": upi_id.strip(),
            "upi_payee_name": upi_payee_name.strip() if upi_payee_name else name.strip(),
            "bank_name": bank_name.strip() if bank_name else "",
            "bank_account_no": bank_account_no.strip() if bank_account_no else "",
            "bank_ifsc": bank_ifsc.strip() if bank_ifsc else "",
            "gst_slabs": gst_slabs_json, "language": language, "master_pin": master_pin.strip()
        })
        supabase_sync.push_to_supabase("workshop_profile", {
            "id": tenant_id, "tenant_id": tenant_id, "client_code": clean_code, "name": name.strip(),
            "owner_name": owner_name.strip(), "phone": whatsapp_mobile.strip(),
            "email": email.strip() if email else "", "address": address.strip(),
            "gstin": gstin.strip() if gstin else "", "logo_url": logo_url.strip() if logo_url else "",
            "upi_id": upi_id.strip(), "upi_payee_name": upi_payee_name.strip() if upi_payee_name else name.strip(),
            "bank_name": bank_name.strip() if bank_name else "",
            "bank_account_no": bank_account_no.strip() if bank_account_no else "",
            "bank_ifsc": bank_ifsc.strip() if bank_ifsc else "",
            "gst_slabs": gst_slabs_json, "language": language, "master_pin": master_pin.strip()
        })
    except Exception as e:
        print(f"[Supabase Client Update Sync] {e}")

    return RedirectResponse(url="/clients?msg=Workshop+configuration+updated+successfully!", status_code=303)

@app.post("/api/admin/toggle-tenant-status")
async def toggle_tenant_status(tenant_id: str = Form(...), new_status: str = Form(...)):
    conn = db.get_db_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE tenants SET status = ? WHERE id = ?", (new_status, tenant_id))
    db.log_audit("TOGGLE_TENANT_STATUS", "tenants", tenant_id, f"Status changed to {new_status}", cursor=cursor)
    conn.commit()
    conn.close()

    try:
        import supabase_sync
        supabase_sync.push_to_supabase("tenants", {"id": tenant_id, "status": new_status})
    except Exception as e:
        print(f"[Supabase Toggle Status] {e}")

    return RedirectResponse(url="/clients?msg=Workshop+status+updated!", status_code=303)

# =========================================================
# 📥 SINGLE WORKSHOP DATA EXPORT (EXCEL & JSON)
# =========================================================
@app.get("/api/admin/clients/{tenant_id}/export-excel")
async def export_client_workshop_excel(tenant_id: str):
    """Generate and download a comprehensive multi-sheet Excel backup workbook for a single workshop."""
    conn = db.get_db_connection()
    cursor = conn.cursor()

    # 1. Fetch Tenant & Profile
    cursor.execute("SELECT * FROM tenants WHERE id = ?", (tenant_id,))
    tenant_row = cursor.fetchone()
    if not tenant_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Client workshop tenant not found")
    tenant = dict(tenant_row)

    cursor.execute("SELECT * FROM workshop_profile WHERE tenant_id = ? OR id = ?", (tenant_id, tenant_id))
    prof_row = cursor.fetchone()
    profile = dict(prof_row) if prof_row else tenant

    # Invoices
    cursor.execute("SELECT * FROM invoices WHERE tenant_id = ? ORDER BY invoice_date DESC, created_at DESC", (tenant_id,))
    invoices = [dict(r) for r in cursor.fetchall()]

    # Job Cards
    cursor.execute("SELECT * FROM job_cards WHERE tenant_id = ? ORDER BY created_at DESC", (tenant_id,))
    jobs = [dict(r) for r in cursor.fetchall()]

    # Inventory
    cursor.execute("SELECT * FROM inventory_items WHERE tenant_id = ? ORDER BY category, part_name", (tenant_id,))
    inventory = [dict(r) for r in cursor.fetchall()]

    # Customers & Vehicles
    cursor.execute("""
    SELECT c.*, v.reg_number as vehicle_reg_no, v.make_model as vehicle_make_model
    FROM customers c
    LEFT JOIN vehicles v ON c.id = v.customer_id
    WHERE c.tenant_id = ?
    ORDER BY c.name ASC
    """, (tenant_id,))
    customers = [dict(r) for r in cursor.fetchall()]

    # Expenses
    cursor.execute("SELECT * FROM expenses WHERE tenant_id = ? ORDER BY expense_date DESC", (tenant_id,))
    expenses = [dict(r) for r in cursor.fetchall()]

    # Purchases
    cursor.execute("SELECT * FROM purchase_bills WHERE tenant_id = ? ORDER BY bill_date DESC", (tenant_id,))
    purchases = [dict(r) for r in cursor.fetchall()]

    # Technicians
    cursor.execute("SELECT * FROM technicians WHERE tenant_id = ? ORDER BY name ASC", (tenant_id,))
    technicians = [dict(r) for r in cursor.fetchall()]

    # Subscriptions
    cursor.execute("SELECT * FROM subscriptions WHERE tenant_id = ? ORDER BY created_at DESC", (tenant_id,))
    subs = [dict(r) for r in cursor.fetchall()]

    db.log_audit("EXPORT_WORKSHOP_DATA_EXCEL", "tenants", tenant_id, f"Exported complete Excel backup for {tenant['name']}", cursor=cursor)
    conn.commit()
    conn.close()

    # Build Excel Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    navy_header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    indigo_fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
    light_blue_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    total_fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")

    font_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, color="E0E7FF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_normal = Font(name="Segoe UI", size=10)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    def style_header(ws, row_idx, num_cols):
        ws.row_dimensions[row_idx].height = 24
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.fill = navy_header_fill
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def auto_fit_columns(ws):
        ws.views.sheetView[0].showGridLines = True
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in (1, 2):
                    continue
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    today_str = datetime.now().strftime("%Y-%m-%d")

    # SHEET 1: SUMMARY
    ws1 = wb.create_sheet(title="Workshop Summary")
    ws1.views.sheetView[0].showGridLines = True
    ws1.merge_cells("A1:D1")
    ws1["A1"] = f"{profile.get('name', tenant.get('name', 'Workshop'))} — DATA BACKUP ARCHIVE"
    ws1["A1"].font = font_title
    ws1["A1"].fill = indigo_fill
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 34

    ws1.merge_cells("A2:D2")
    ws1["A2"] = f"AutoVolt Cloud SaaS Vendor Platform | Tenant ID: {tenant_id} | Exported Date: {today_str}"
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = navy_header_fill
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    ws1.append([])

    meta_fields = [
        ("Workshop Legal Name", profile.get("name", "")),
        ("Owner / Proprietor", profile.get("owner_name", "")),
        ("Primary Phone", profile.get("phone", "")),
        ("WhatsApp Number", tenant.get("whatsapp_mobile", profile.get("phone", ""))),
        ("Email Address", profile.get("email", "") or "N/A"),
        ("Workshop Address", profile.get("address", "")),
        ("GSTIN", profile.get("gstin", "") or "Unregistered / N/A"),
        ("UPI VPA (Payee ID)", profile.get("upi_id", "")),
        ("Bank Name", profile.get("bank_name", "") or "N/A"),
        ("Bank Account Number", profile.get("bank_account_no", "") or "N/A"),
        ("Bank IFSC Code", profile.get("bank_ifsc", "") or "N/A"),
        ("Bank Branch", profile.get("bank_branch", "") or "N/A"),
        ("License Status", tenant.get("status", "ACTIVE")),
        ("Subscription Plan", tenant.get("subscription_plan", "PRO_YEARLY")),
        ("License Valid Until", tenant.get("subscription_end_date", "N/A")),
        ("Master Security PIN", profile.get("master_pin", "1234")),
        ("Total Invoices Generated", len(invoices)),
        ("Total Revenue (₹)", sum(float(i.get("amount_paid") or 0.0) for i in invoices)),
        ("Total Khata Dues (₹)", sum(float(i.get("balance_due") or 0.0) for i in invoices)),
        ("Total Job Cards", len(jobs)),
        ("Total Inventory Parts", len(inventory)),
        ("Total Registered Customers", len(customers)),
        ("Total Staff / Technicians", len(technicians)),
    ]

    ws1.append(["Field Property", "Configured Value", "", ""])
    ws1.row_dimensions[4].height = 22
    for c in range(1, 3):
        ws1.cell(row=4, column=c).font = font_header
        ws1.cell(row=4, column=c).fill = navy_header_fill

    curr_r = 5
    for field, val in meta_fields:
        ws1.append([field, val])
        ws1.cell(row=curr_r, column=1).font = font_bold
        ws1.cell(row=curr_r, column=1).border = thin_border
        ws1.cell(row=curr_r, column=1).fill = light_blue_fill
        
        c2 = ws1.cell(row=curr_r, column=2)
        c2.font = font_normal
        c2.border = thin_border
        if "₹" in field:
            c2.number_format = '₹#,##0.00'
            c2.alignment = Alignment(horizontal="right")
        curr_r += 1

    ws1.column_dimensions['A'].width = 28
    ws1.column_dimensions['B'].width = 45

    # SHEET 2: INVOICES
    ws2 = wb.create_sheet(title="Invoices & Billing")
    headers2 = [
        "Sl No", "Invoice No", "Date", "Type", "Job Number", "Customer Name", "Phone",
        "Vehicle Reg No", "Make & Model", "Labor (₹)", "Parts (₹)", "Taxable Subtotal (₹)",
        "CGST (₹)", "SGST (₹)", "Total Tax (₹)", "Discount (₹)", "Grand Total (₹)",
        "Amount Paid (₹)", "Balance Due (₹)", "Payment Status", "Payment Mode"
    ]
    ws2.append(headers2)
    style_header(ws2, 1, len(headers2))

    tot_taxable = 0.0
    tot_grand = 0.0
    tot_paid = 0.0
    tot_due = 0.0

    for idx, inv in enumerate(invoices, 1):
        taxable = float(inv.get("taxable_subtotal") or 0.0)
        grand = float(inv.get("grand_total") or 0.0)
        paid = float(inv.get("amount_paid") or 0.0)
        due = float(inv.get("balance_due") or 0.0)
        tot_taxable += taxable
        tot_grand += grand
        tot_paid += paid
        tot_due += due

        row = [
            idx, inv.get("invoice_number", ""), inv.get("invoice_date", ""), inv.get("invoice_type", ""),
            inv.get("job_number", ""), inv.get("customer_name", ""), inv.get("customer_phone", ""),
            inv.get("vehicle_reg_no", ""), inv.get("vehicle_make_model", ""),
            float(inv.get("subtotal_labor") or 0.0), float(inv.get("subtotal_parts") or 0.0), taxable,
            float(inv.get("cgst_total") or 0.0), float(inv.get("sgst_total") or 0.0), float(inv.get("tax_total") or 0.0),
            float(inv.get("discount_amount") or 0.0), grand, paid, due,
            inv.get("payment_status", "PAID"), inv.get("payment_mode", "UPI")
        ]
        ws2.append(row)
        r_num = idx + 1
        for col_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=r_num, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx in range(10, 20):
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 3, 20, 21):
                cell.alignment = Alignment(horizontal="center")

    if invoices:
        tot_r_num = len(invoices) + 2
        ws2.append([
            "TOTAL", "", "", "", "", "", "", "", "",
            "", "", tot_taxable, "", "", "", "", tot_grand, tot_paid, tot_due, "", ""
        ])
        for c in range(1, len(headers2) + 1):
            cell = ws2.cell(row=tot_r_num, column=c)
            cell.font = font_bold
            cell.fill = total_fill
            cell.border = thin_border
            if c in (12, 17, 18, 19):
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")

    auto_fit_columns(ws2)

    # SHEET 3: JOB CARDS
    ws3 = wb.create_sheet(title="Job Cards")
    headers3 = [
        "Sl No", "Job Card No", "Status", "Vehicle Reg", "Make & Model",
        "Customer Name", "Customer Phone", "Odometer (km)", "Assigned Tech",
        "Complaints / Issues", "Labor Total (₹)", "Parts Total (₹)", "Job Total (₹)", "Created Date", "Delivered Date"
    ]
    ws3.append(headers3)
    style_header(ws3, 1, len(headers3))

    for idx, j in enumerate(jobs, 1):
        comp = j.get("complaints", "")
        try:
            c_list = json.loads(comp) if comp else []
            comp_str = ", ".join(c_list) if isinstance(c_list, list) else str(comp)
        except Exception:
            comp_str = str(comp)

        ws3.append([
            idx, j.get("job_number", ""), j.get("status", "RECEIVED"), j.get("vehicle_reg_no", ""),
            j.get("vehicle_make_model", ""), j.get("customer_name", ""), j.get("customer_phone", ""),
            j.get("odometer", 0), j.get("assigned_technician_name", "") or "Unassigned",
            comp_str, float(j.get("total_labor") or 0.0), float(j.get("total_parts") or 0.0),
            float(j.get("grand_total") or 0.0), j.get("created_at", "")[:10], j.get("delivered_at", "") or "In Progress"
        ])
        r_num = idx + 1
        for col_idx in range(1, len(headers3) + 1):
            cell = ws3.cell(row=r_num, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx in (11, 12, 13):
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 3, 14, 15):
                cell.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws3)

    # SHEET 4: INVENTORY
    ws4 = wb.create_sheet(title="Inventory & Stock")
    headers4 = [
        "Sl No", "Part Name", "Part Number", "Barcode / SKU", "Category",
        "Stock Qty", "Min Alert Qty", "Cost Price (₹)", "Selling Price (₹)", "Tax Rate (%)",
        "HSN Code", "Location Rack", "Position Bin", "Total Stock Value (₹)"
    ]
    ws4.append(headers4)
    style_header(ws4, 1, len(headers4))

    for idx, item in enumerate(inventory, 1):
        qty = int(item.get("stock_qty") or 0)
        cp = float(item.get("cost_price") or 0.0)
        sp = float(item.get("selling_price") or 0.0)
        stk_val = qty * cp

        ws4.append([
            idx, item.get("part_name", ""), item.get("part_number", ""), item.get("barcode", ""),
            item.get("category", ""), qty, item.get("min_stock_alert", 5),
            cp, sp, f"{item.get('tax_rate', 18.0)}%",
            item.get("hsn_code", "8536"), item.get("location_rack", "Main"), item.get("position_bin", ""), stk_val
        ])
        r_num = idx + 1
        for col_idx in range(1, len(headers4) + 1):
            cell = ws4.cell(row=r_num, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx in (8, 9, 14):
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 6, 7, 10, 11):
                cell.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws4)

    # SHEET 5: CUSTOMERS & KHATA
    ws5 = wb.create_sheet(title="Customers & Khata")
    headers5 = [
        "Sl No", "Customer Name", "Phone", "Vehicle Reg", "Make & Model",
        "Total Visits", "Khata Dues (₹)", "Address", "GSTIN"
    ]
    ws5.append(headers5)
    style_header(ws5, 1, len(headers5))

    for idx, c in enumerate(customers, 1):
        dues = float(c.get("khata_balance") or 0.0)
        ws5.append([
            idx, c.get("name", ""), c.get("phone", ""), c.get("vehicle_reg_no", "") or "N/A",
            c.get("vehicle_make_model", "") or "General", c.get("total_visits", 1),
            dues, c.get("address", "") or "N/A", c.get("gstin", "") or "Unregistered"
        ])
        r_num = idx + 1
        for col_idx in range(1, len(headers5) + 1):
            cell = ws5.cell(row=r_num, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx == 7:
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 6):
                cell.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws5)

    # SHEET 6: PURCHASES
    ws6 = wb.create_sheet(title="Purchases")
    headers6 = [
        "Sl No", "Bill Number", "Bill Date", "Supplier Name", "Total Amount (₹)",
        "CGST (₹)", "SGST (₹)", "IGST (₹)", "Payment Mode", "Status", "Notes"
    ]
    ws6.append(headers6)
    style_header(ws6, 1, len(headers6))

    for idx, p in enumerate(purchases, 1):
        ws6.append([
            idx, p.get("bill_number", ""), p.get("bill_date", ""), p.get("supplier_name", ""),
            float(p.get("total_amount") or 0.0), float(p.get("cgst_amount") or 0.0),
            float(p.get("sgst_amount") or 0.0), float(p.get("igst_amount") or 0.0),
            p.get("payment_mode", "BANK"), p.get("status", "CONFIRMED"), p.get("notes", "") or ""
        ])
        r_num = idx + 1
        for col_idx in range(1, len(headers6) + 1):
            cell = ws6.cell(row=r_num, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx in (5, 6, 7, 8):
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 3, 9, 10):
                cell.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws6)

    # SHEET 7: EXPENSES
    ws7 = wb.create_sheet(title="Expenses")
    headers7 = [
        "Sl No", "Date", "Expense Title", "Category", "Amount (₹)",
        "Payment Mode", "Paid To", "Reference No"
    ]
    ws7.append(headers7)
    style_header(ws7, 1, len(headers7))

    for idx, exp in enumerate(expenses, 1):
        ws7.append([
            idx, exp.get("expense_date", ""), exp.get("title", ""), exp.get("category", ""),
            float(exp.get("amount") or 0.0), exp.get("payment_mode", "CASH"),
            exp.get("paid_to", "") or "", exp.get("reference_no", "") or ""
        ])
        r_num = idx + 1
        for col_idx in range(1, len(headers7) + 1):
            cell = ws7.cell(row=r_num, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx == 5:
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 2, 6):
                cell.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws7)

    # SHEET 8: TECHNICIANS
    ws8 = wb.create_sheet(title="Staff & Technicians")
    headers8 = [
        "Sl No", "Name", "Phone", "Role", "Commission (%)",
        "Daily Wage (₹)", "Advance Balance (₹)", "Total Jobs Handled", "Status"
    ]
    ws8.append(headers8)
    style_header(ws8, 1, len(headers8))

    for idx, tech in enumerate(technicians, 1):
        ws8.append([
            idx, tech.get("name", ""), tech.get("phone", ""), tech.get("role", ""),
            float(tech.get("commission_pct") or 0.0), float(tech.get("daily_wage") or 0.0),
            float(tech.get("current_advance_balance") or 0.0), tech.get("total_jobs", 0),
            tech.get("status", "Active")
        ])
        r_num = idx + 1
        for col_idx in range(1, len(headers8) + 1):
            cell = ws8.cell(row=r_num, column=col_idx)
            cell.font = font_normal
            cell.border = thin_border
            if col_idx in (6, 7):
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 5, 8, 9):
                cell.alignment = Alignment(horizontal="center")

    auto_fit_columns(ws8)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    clean_name = "".join(c for c in tenant.get("name", "Workshop") if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
    filename = f"{clean_name}_Backup_{today_str}.xlsx"

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@app.get("/api/admin/clients/{tenant_id}/export-json")
async def export_client_workshop_json(tenant_id: str):
    """Download full JSON dump of a single workshop's data for system backup or migration."""
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM tenants WHERE id = ?", (tenant_id,))
    tenant_row = cursor.fetchone()
    if not tenant_row:
        conn.close()
        raise HTTPException(status_code=404, detail="Client workshop tenant not found")
    tenant = dict(tenant_row)

    cursor.execute("SELECT * FROM workshop_profile WHERE tenant_id = ? OR id = ?", (tenant_id, tenant_id))
    prof_row = cursor.fetchone()
    profile = dict(prof_row) if prof_row else tenant

    cursor.execute("SELECT * FROM invoices WHERE tenant_id = ? ORDER BY invoice_date DESC", (tenant_id,))
    invoices = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM job_cards WHERE tenant_id = ? ORDER BY created_at DESC", (tenant_id,))
    jobs = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM inventory_items WHERE tenant_id = ?", (tenant_id,))
    inventory = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM customers WHERE tenant_id = ?", (tenant_id,))
    customers = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM vehicles WHERE tenant_id = ?", (tenant_id,))
    vehicles = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM expenses WHERE tenant_id = ?", (tenant_id,))
    expenses = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM purchase_bills WHERE tenant_id = ?", (tenant_id,))
    purchases = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM technicians WHERE tenant_id = ?", (tenant_id,))
    technicians = [dict(r) for r in cursor.fetchall()]

    cursor.execute("SELECT * FROM subscriptions WHERE tenant_id = ?", (tenant_id,))
    subs = [dict(r) for r in cursor.fetchall()]

    db.log_audit("EXPORT_WORKSHOP_DATA_JSON", "tenants", tenant_id, f"Exported complete JSON backup for {tenant['name']}", cursor=cursor)
    conn.commit()
    conn.close()

    payload = {
        "export_date": datetime.now().isoformat(),
        "tenant": tenant,
        "workshop_profile": profile,
        "subscriptions": subs,
        "invoices": invoices,
        "job_cards": jobs,
        "inventory": inventory,
        "customers": customers,
        "vehicles": vehicles,
        "expenses": expenses,
        "purchases": purchases,
        "technicians": technicians
    }

    clean_name = "".join(c for c in tenant.get("name", "Workshop") if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
    filename = f"{clean_name}_Backup_{datetime.now().strftime('%Y-%m-%d')}.json"

    return Response(
        content=json.dumps(payload, indent=2, default=str),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

# =========================================================
# 👥 WORKSHOP USERS & OTP MANAGEMENT
# =========================================================
@app.get("/users", response_class=HTMLResponse)
async def users_page(request: Request, msg: Optional[str] = None):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT u.*, t.name as tenant_name 
        FROM users u 
        LEFT JOIN tenants t ON u.tenant_id = t.id 
        ORDER BY u.created_at DESC
    """)
    users = [dict(row) for row in cursor.fetchall()]

    cursor.execute("SELECT id, name, owner_name FROM tenants ORDER BY name ASC")
    tenants = [dict(row) for row in cursor.fetchall()]

    conn.close()

    return templates.TemplateResponse(
        request=request,
        name="admin_hq/users.html",
        context={
            "active_page": "users",
            "users": users,
            "tenants": tenants,
            "msg": msg
        }
    )

@app.post("/api/admin/users/create")
async def create_user_account(
    tenant_id: str = Form(...),
    full_name: str = Form(...),
    username: str = Form(...),
    whatsapp_mobile: str = Form(...),
    role: str = Form("STAFF"),
    password: str = Form(...),
    permissions: Optional[List[str]] = Form(None)
):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT COUNT(*) FROM users WHERE username = ?", (username.strip(),))
    if cursor.fetchone()[0] > 0:
        conn.close()
        return RedirectResponse(url="/users?msg=Error:+Username+already+exists.+Please+choose+another.", status_code=303)

    user_id = f"user_{uuid.uuid4().hex[:8]}"
    pw_hash = hash_password(password.strip())
    perms_str = ",".join(permissions) if permissions else ""
    if not perms_str:
        if role == "CASHIER":
            perms_str = "can_bill,can_jobs,can_khata"
        elif role == "STOREKEEPER":
            perms_str = "can_inventory,can_purchases"
        elif role in ["ADMIN", "OWNER"]:
            perms_str = "can_bill,can_inventory,can_purchases,can_wages,can_reports,can_settings,can_jobs,can_khata"

    cursor.execute("""
    INSERT INTO users (id, username, password_hash, full_name, role, whatsapp_mobile, phone, tenant_id, permissions, status, created_at)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'ACTIVE', datetime('now'))
    """, (
        user_id, username.strip(), pw_hash, full_name.strip(), role,
        whatsapp_mobile.strip(), whatsapp_mobile.strip(), tenant_id, perms_str
    ))

    db.log_audit("CREATE_USER", "users", user_id, f"Created user {username} ({role}) for tenant {tenant_id}", cursor=cursor)
    conn.commit()
    conn.close()

    return RedirectResponse(url="/users?msg=User+account+created+successfully!", status_code=303)

@app.post("/api/admin/users/reset-password")
async def reset_user_password(
    user_id: str = Form(...),
    username: str = Form(...),
    new_password: str = Form(...)
):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    pw_hash = hash_password(new_password.strip())
    cursor.execute("""
    UPDATE users 
    SET password_hash = ?, otp_code = NULL, otp_expiry = NULL, reset_token = NULL, reset_token_expiry = NULL 
    WHERE id = ?
    """, (pw_hash, user_id))

    db.log_audit("ADMIN_RESET_PASSWORD", "users", user_id, f"Master Admin reset password for {username}", cursor=cursor)
    conn.commit()
    conn.close()

    return RedirectResponse(url="/users?msg=Password+reset+successfully+for+user+@" + username, status_code=303)

# =========================================================
# 🔑 SAAS LICENSES & SUBSCRIPTIONS
# =========================================================
@app.get("/subscriptions", response_class=HTMLResponse)
async def subscriptions_page(request: Request, tenant_id: Optional[str] = None, msg: Optional[str] = None):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    query = """
        SELECT s.*, t.name as tenant_name 
        FROM subscriptions s 
        LEFT JOIN tenants t ON s.tenant_id = t.id 
    """
    params = []
    if tenant_id:
        query += " WHERE s.tenant_id = ? "
        params.append(tenant_id)
    query += " ORDER BY s.created_at DESC"

    cursor.execute(query, tuple(params))
    subscriptions = [dict(row) for row in cursor.fetchall()]

    cursor.execute("SELECT id, name, subscription_end_date FROM tenants ORDER BY name ASC")
    tenants = [dict(row) for row in cursor.fetchall()]

    conn.close()

    return templates.TemplateResponse(
        request=request,
        name="admin_hq/subscriptions.html",
        context={
            "active_page": "subscriptions",
            "subscriptions": subscriptions,
            "tenants": tenants,
            "selected_tenant_id": tenant_id,
            "msg": msg
        }
    )

@app.post("/api/admin/subscriptions/renew")
async def renew_subscription(
    tenant_id: str = Form(...),
    billing_cycle: str = Form("YEARLY"),
    payment_mode: str = Form("UPI"),
    taxable_amount: float = Form(9999.0),
    payment_ref: Optional[str] = Form(None)
):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM tenants WHERE id = ?", (tenant_id,))
    tenant = cursor.fetchone()
    if not tenant:
        conn.close()
        return RedirectResponse(url="/subscriptions?msg=Tenant+not+found", status_code=303)

    tenant_name = tenant["name"]
    current_end = tenant["subscription_end_date"]
    today = datetime.now().date()

    # If already active in future, extend from current end date; otherwise start from today
    start_dt = today
    if current_end:
        try:
            curr_dt = datetime.strptime(current_end, "%Y-%m-%d").date()
            if curr_dt > today:
                start_dt = curr_dt
        except Exception:
            pass

    days = 365 if billing_cycle == "YEARLY" else 30
    new_end_dt = start_dt + timedelta(days=days)

    start_str = start_dt.strftime("%Y-%m-%d")
    end_str = new_end_dt.strftime("%Y-%m-%d")

    cgst = round(taxable_amount * 0.09, 2)
    sgst = round(taxable_amount * 0.09, 2)
    total_amount = round(taxable_amount + cgst + sgst, 2)

    # Update Tenant
    cursor.execute("""
    UPDATE tenants SET 
        status = 'ACTIVE', 
        subscription_plan = ?,
        subscription_start_date = ?,
        subscription_end_date = ?,
        price_paid = price_paid + ?
    WHERE id = ?
    """, (f"PRO_{billing_cycle}", start_str, end_str, total_amount, tenant_id))

    # Insert Subscription Record
    sub_id = f"sub_{uuid.uuid4().hex[:8]}"
    cursor.execute("""
    INSERT INTO subscriptions (id, tenant_id, plan_name, start_date, end_date, is_active, billing_cycle, price_paid, notes, created_at)
    VALUES (?, ?, ?, ?, ?, 1, ?, ?, 'License Renewal', datetime('now'))
    """, (sub_id, tenant_id, f"PRO_{billing_cycle}", start_str, end_str, billing_cycle, total_amount))

    # Auto-generate SaaS Vendor Invoice for Vendor Books
    inv_id = f"sinv_{uuid.uuid4().hex[:8]}"
    cursor.execute("SELECT COUNT(*) FROM saas_vendor_invoices")
    inv_count = cursor.fetchone()[0] + 1
    inv_no = f"HQ-INV-2026-{inv_count:03d}"
    cursor.execute("SELECT id FROM saas_vendor_invoices WHERE invoice_no = ?", (inv_no,))
    while cursor.fetchone():
        inv_count += 1
        inv_no = f"HQ-INV-2026-{inv_count:03d}"
        cursor.execute("SELECT id FROM saas_vendor_invoices WHERE invoice_no = ?", (inv_no,))

    cursor.execute("""
    INSERT INTO saas_vendor_invoices (
        id, invoice_no, tenant_id, tenant_name, invoice_date, plan_type,
        subscription_start, subscription_end, taxable_amount, cgst_rate, cgst_amount,
        sgst_rate, sgst_amount, total_amount, payment_status, payment_mode, payment_ref, notes, created_at
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 9.0, ?, 9.0, ?, ?, 'PAID', ?, ?, 'SaaS License Renewal', datetime('now'))
    """, (
        inv_id, inv_no, tenant_id, tenant_name, datetime.now().strftime("%Y-%m-%d"), billing_cycle,
        start_str, end_str, taxable_amount, cgst, sgst, total_amount, payment_mode, payment_ref or "UPI-PAID"
    ))

    db.log_audit("RENEW_SUBSCRIPTION", "subscriptions", sub_id, f"Renewed license for {tenant_name} until {end_str}", cursor=cursor)
    conn.commit()
    conn.close()

    try:
        import supabase_sync
        supabase_sync.push_to_supabase("tenants", {
            "id": tenant_id, "status": "ACTIVE", "subscription_plan": f"PRO_{billing_cycle}",
            "subscription_start_date": start_str, "subscription_end_date": end_str
        })
        supabase_sync.push_to_supabase("subscriptions", {
            "id": sub_id, "tenant_id": tenant_id, "plan_name": f"PRO_{billing_cycle}",
            "start_date": start_str, "end_date": end_str, "is_active": 1,
            "billing_cycle": billing_cycle, "price_paid": total_amount, "notes": "License Renewal"
        })
        supabase_sync.push_to_supabase("saas_vendor_invoices", {
            "id": inv_id, "invoice_no": inv_no, "tenant_id": tenant_id, "tenant_name": tenant_name,
            "invoice_date": datetime.now().strftime("%Y-%m-%d"), "plan_type": billing_cycle,
            "subscription_start": start_str, "subscription_end": end_str,
            "taxable_amount": taxable_amount, "cgst_rate": 9.0, "cgst_amount": cgst,
            "sgst_rate": 9.0, "sgst_amount": sgst, "total_amount": total_amount,
            "payment_status": "PAID", "payment_mode": payment_mode, "payment_ref": payment_ref or "UPI-PAID",
            "notes": "SaaS License Renewal"
        })
    except Exception as e:
        print(f"[Supabase Renewal Sync] {e}")

    return RedirectResponse(url=f"/subscriptions?tenant_id={tenant_id}&msg=License+renewed+until+{end_str}+and+invoice+generated!", status_code=303)

# =========================================================
# 💰 SAAS VENDOR ACCOUNTING, REVENUE & GSTR-1 AUDIT
# =========================================================
@app.get("/accounting", response_class=HTMLResponse)
async def accounting_page(request: Request, fy: Optional[str] = "2025-2026"):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    query = "SELECT * FROM saas_vendor_invoices"
    params = []

    if fy == "2025-2026":
        query += " WHERE invoice_date >= '2025-04-01' AND invoice_date <= '2026-03-31'"
    elif fy == "2026-2027":
        query += " WHERE invoice_date >= '2026-04-01' AND invoice_date <= '2027-03-31'"

    query += " ORDER BY invoice_date DESC"
    cursor.execute(query, tuple(params))
    invoices = [dict(row) for row in cursor.fetchall()]

    taxable_total = sum(i.get("taxable_amount", 0.0) for i in invoices)
    cgst_total = sum(i.get("cgst_amount", 0.0) for i in invoices)
    sgst_total = sum(i.get("sgst_amount", 0.0) for i in invoices)
    gross_total = sum(i.get("total_amount", 0.0) for i in invoices)

    summary = {
        "taxable_total": taxable_total,
        "cgst_total": cgst_total,
        "sgst_total": sgst_total,
        "gross_total": gross_total,
        "invoices_count": len(invoices)
    }

    conn.close()

    return templates.TemplateResponse(
        request=request,
        name="admin_hq/accounting.html",
        context={
            "active_page": "accounting",
            "invoices": invoices,
            "summary": summary,
            "current_fy": fy
        }
    )

@app.get("/api/admin/accounting/export-gstr1-excel")
async def export_vendor_gstr1_excel(fy: Optional[str] = "2025-2026"):
    """Export SaaS Vendor GSTR-1 Sales Report in styled Excel workbook for CA"""
    conn = db.get_db_connection()
    cursor = conn.cursor()

    query = "SELECT * FROM saas_vendor_invoices"
    if fy == "2025-2026":
        query += " WHERE invoice_date >= '2025-04-01' AND invoice_date <= '2026-03-31'"
    elif fy == "2026-2027":
        query += " WHERE invoice_date >= '2026-04-01' AND invoice_date <= '2027-03-31'"
    query += " ORDER BY invoice_date ASC"

    cursor.execute(query)
    invoices = [dict(row) for row in cursor.fetchall()]

    cursor.execute("SELECT * FROM saas_vendor_profile LIMIT 1")
    prof_row = cursor.fetchone()
    vendor = dict(prof_row) if prof_row else {
        "company_name": "AutoVolt Technologies Pvt Ltd",
        "gstin": "32AABCA9876C1Z5"
    }

    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GSTR-1 SaaS Sales"
    ws.views.sheetView[0].showGridLines = True

    # Styling colors
    purple_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")
    dark_header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    title_font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    sub_font = Font(name="Segoe UI", size=10, color="E0E7FF")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    normal_font = Font(name="Segoe UI", size=10)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title Banner
    ws.merge_cells("A1:J1")
    ws["A1"] = f"{vendor.get('company_name', 'AutoVolt Technologies')} — GSTR-1 SAAS SALES AUDIT REPORT"
    ws["A1"].font = title_font
    ws["A1"].fill = purple_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    ws["A2"] = f"GSTIN: {vendor.get('gstin', '')} | Financial Year: {fy} | SAC: 997331 (Software Licensing & SaaS Subscriptions)"
    ws["A2"].font = sub_font
    ws["A2"].fill = dark_header_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        "Sl No", "Invoice No", "Invoice Date", "Client Garage Name", "SAC Code",
        "Taxable Amount (₹)", "CGST (9%)", "SGST (9%)", "Total Tax (18%)", "Invoice Total (₹)"
    ]

    ws.append([]) # Row 3 blank
    ws.append(headers) # Row 4
    ws.row_dimensions[4].height = 24

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = header_font
        cell.fill = dark_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_taxable = 0.0
    total_cgst = 0.0
    total_sgst = 0.0
    total_gross = 0.0

    current_row = 5
    for idx, inv in enumerate(invoices, 1):
        taxable = inv.get("taxable_amount", 0.0)
        cgst = inv.get("cgst_amount", 0.0)
        sgst = inv.get("sgst_amount", 0.0)
        tot_tax = cgst + sgst
        inv_tot = inv.get("total_amount", 0.0)

        total_taxable += taxable
        total_cgst += cgst
        total_sgst += sgst
        total_gross += inv_tot

        row_data = [
            idx,
            inv.get("invoice_no", ""),
            inv.get("invoice_date", ""),
            inv.get("tenant_name", ""),
            "997331",
            taxable,
            cgst,
            sgst,
            tot_tax,
            inv_tot
        ]
        ws.append(row_data)

        for c_idx in range(1, 11):
            c = ws.cell(row=current_row, column=c_idx)
            c.font = normal_font
            c.border = thin_border
            if c_idx in (6, 7, 8, 9, 10):
                c.number_format = '₹#,##0.00'
                c.alignment = Alignment(horizontal="right")
            elif c_idx in (1, 3, 5):
                c.alignment = Alignment(horizontal="center")

        current_row += 1

    # Total Summary Row
    summary_row = [
        "TOTAL", "", "", "", "",
        total_taxable, total_cgst, total_sgst, (total_cgst + total_sgst), total_gross
    ]
    ws.append(summary_row)
    ws.row_dimensions[current_row].height = 24

    for c_idx in range(1, 11):
        c = ws.cell(row=current_row, column=c_idx)
        c.font = bold_font
        c.border = thin_border
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        if c_idx in (6, 7, 8, 9, 10):
            c.number_format = '₹#,##0.00'
            c.alignment = Alignment(horizontal="right")

    # Column Widths
    col_widths = [8, 18, 14, 30, 12, 20, 15, 15, 18, 20]
    for idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"GSTR1_SaaS_Vendor_Audit_{fy}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# =========================================================
# ⚙️ SAAS VENDOR COMPANY SETTINGS
# =========================================================
@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request, msg: Optional[str] = None):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM saas_vendor_profile LIMIT 1")
    prof = cursor.fetchone()
    profile = dict(prof) if prof else {}

    conn.close()

    return templates.TemplateResponse(
        request=request,
        name="admin_hq/settings.html",
        context={
            "active_page": "settings",
            "profile": profile,
            "msg": msg
        }
    )

@app.post("/api/admin/vendor-settings/update")
async def update_vendor_settings(
    company_name: str = Form(...),
    tagline: Optional[str] = Form(None),
    gstin: str = Form(...),
    email: str = Form(...),
    phone: Optional[str] = Form(None),
    whatsapp_phone: Optional[str] = Form(None),
    address: Optional[str] = Form(None),
    upi_id: str = Form(...),
    upi_payee_name: Optional[str] = Form(None),
    bank_name: Optional[str] = Form(None),
    bank_account_no: Optional[str] = Form(None),
    bank_ifsc: Optional[str] = Form(None),
    monthly_fee: float = Form(999.0),
    yearly_fee: float = Form(9999.0)
):
    conn = db.get_db_connection()
    cursor = conn.cursor()

    cursor.execute("""
    UPDATE saas_vendor_profile SET
        company_name = ?, tagline = ?, gstin = ?, email = ?, phone = ?, whatsapp_phone = ?,
        address = ?, upi_id = ?, upi_payee_name = ?, bank_name = ?, bank_account_no = ?, bank_ifsc = ?,
        monthly_fee = ?, yearly_fee = ?, updated_at = datetime('now')
    WHERE id = 'vendor_main'
    """, (
        company_name.strip(), tagline.strip() if tagline else "", gstin.strip(), email.strip(),
        phone.strip() if phone else "", whatsapp_phone.strip() if whatsapp_phone else "",
        address.strip() if address else "", upi_id.strip(), upi_payee_name.strip() if upi_payee_name else company_name.strip(),
        bank_name.strip() if bank_name else "", bank_account_no.strip() if bank_account_no else "",
        bank_ifsc.strip() if bank_ifsc else "", monthly_fee, yearly_fee
    ))

    db.log_audit("UPDATE_VENDOR_PROFILE", "saas_vendor_profile", "vendor_main", f"Updated SaaS company profile", cursor=cursor)
    conn.commit()
    conn.close()

    return RedirectResponse(url="/settings?msg=SaaS+vendor+settings+saved+successfully!", status_code=303)

if __name__ == "__main__":
    import uvicorn
    print("🚀 Starting AutoVolt Cloud HQ (SaaS Master Vendor Portal) on http://localhost:8500 ...")
    uvicorn.run("saas_admin_app:app", host="0.0.0.0", port=8500, reload=True)
