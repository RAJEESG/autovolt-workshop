import io
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Tuple

def generate_inventory_import_template() -> bytes:
    """Generate professional Excel template (.xlsx) for bulk opening inventory upload"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory_Template"

    # Header styling
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")

    headers = [
        "Part Name *", "Part Number (OEM / Mfr No)", "Barcode (Box Barcode) *", "Category *", "Opening Stock Qty *",
        "Cost Price (₹ Base) *", "Selling Price (₹ Base) *", "GST Tax %", "HSN Code",
        "Unit", "Warranty Months", "Location / Rack", "Shelf / Bin Position"
    ]
    ws.append(headers)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Sample rows
    sample_rows = [
        ["Lucas 4-Pin Horn Relay 12V 30A", "LUC-REL-4P", "89012340001", "Relays & Fuses", 40, 85.00, 140.00, 18.0, "8536", "pcs", 6, "Rack A1", "Shelf 2 / Bin 4"],
        ["Bosch 5-Pin Changeover Relay 12V", "BOS-REL-5P", "89012340002", "Relays & Fuses", 25, 110.00, 180.00, 18.0, "8536", "pcs", 6, "Rack A1", "Shelf 2 / Bin 5"],
        ["Blade Fuse 15A Blue Standard", "FUSE-15A", "89012340003", "Relays & Fuses", 100, 3.50, 10.00, 18.0, "8536", "pcs", 0, "Rack A2", "Drawer 1"],
        ["Philips H4 12V 100/90W Rally Bulb", "PHI-H4-100", "89012340005", "Bulbs & LEDs", 20, 145.00, 220.00, 18.0, "8539", "pcs", 6, "Rack B1", "Top Shelf"],
        ["Amaron Hi-Way 12V 35Ah Battery", "AAM-HW-35", "89012340011", "Batteries", 8, 3400.00, 4250.00, 28.0, "8507", "pcs", 36, "Battery Bay", "Floor Pallet"],
        ["Universal Reverse Parking Camera", "REV-CAM-01", "89012340016", "Accessories", 10, 480.00, 850.00, 18.0, "8528", "pcs", 12, "Rack E1", "Bin 12"]
    ]

    for row_data in sample_rows:
        ws.append(row_data)

    # Set column widths
    column_widths = [36, 26, 24, 22, 18, 18, 18, 12, 12, 10, 16, 18, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def _clean_header(col_name: str) -> str:
    cleaned = str(col_name).strip().replace("*", "").replace("₹", "").replace("%", "").strip()
    return "_".join(re.split(r'[\s/_()]+', cleaned.lower())).strip("_")

def parse_inventory_file(file_bytes: bytes, filename: str) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Parse uploaded Excel (.xlsx) or CSV file and return list of valid items and error list"""
    items = []
    errors = []

    try:
        if filename.lower().endswith(".csv"):
            df = pd.read_csv(io.BytesIO(file_bytes))
        else:
            df = pd.read_excel(io.BytesIO(file_bytes))
    except Exception as e:
        return [], [f"Failed to read file: {e}"]

    # Standardize column headers
    df.columns = [_clean_header(c) for c in df.columns]

    for index, row in df.iterrows():
        row_num = index + 2
        part_name = str(row.get("part_name") or row.get("name") or row.get("part") or row.get("item") or "").strip()
        barcode = str(row.get("barcode_box_barcode") or row.get("barcode") or row.get("sku") or row.get("box_barcode") or "").strip()

        # Handle numeric float barcodes from Excel e.g. 89012340001.0
        if barcode.endswith(".0"):
            barcode = barcode[:-2]

        if not part_name or part_name.lower() == "nan":
            continue
        if not barcode or barcode.lower() == "nan":
            errors.append(f"Row {row_num}: Missing barcode for '{part_name}'")
            continue

        try:
            part_number = str(row.get("part_number_oem_mfr_no") or row.get("part_number") or row.get("part_no") or row.get("oem_no") or "").strip().upper()
            if part_number.lower() == "nan":
                part_number = ""

            category = str(row.get("category") or "Relays & Fuses").strip()
            stock_qty = int(float(row.get("opening_stock_qty") or row.get("stock_qty") or row.get("qty") or row.get("stock") or 0))
            cost_price = float(row.get("cost_price_base") or row.get("cost_price") or row.get("cost") or row.get("purchase_price") or 0.0)
            selling_price = float(row.get("selling_price_base") or row.get("selling_price") or row.get("price") or row.get("selling_rate") or row.get("mrp") or 0.0)
            tax_rate = float(row.get("gst_tax") or row.get("tax_rate") or row.get("gst") or row.get("tax") or 18.0)
            hsn_code = str(row.get("hsn_code") or row.get("hsn") or "8536").strip()
            if hsn_code.endswith(".0"):
                hsn_code = hsn_code[:-2]
            unit = str(row.get("unit") or "pcs").strip()
            warranty_months = int(float(row.get("warranty_months") or row.get("warranty") or 0))
            location_rack = str(row.get("location_rack") or row.get("rack") or row.get("location") or "Main").strip()
            if location_rack.lower() == "nan":
                location_rack = "Main"
            position_bin = str(row.get("shelf_bin_position") or row.get("position_bin") or row.get("bin") or row.get("shelf") or row.get("position") or "").strip()
            if position_bin.lower() == "nan":
                position_bin = ""

            items.append({
                "part_name": part_name,
                "part_number": part_number,
                "barcode": barcode,
                "sku": barcode,
                "category": category if category.lower() != "nan" else "Relays & Fuses",
                "stock_qty": stock_qty,
                "cost_price": cost_price,
                "selling_price": selling_price,
                "tax_rate": tax_rate,
                "hsn_code": hsn_code if hsn_code.lower() != "nan" else "8536",
                "unit": unit if unit.lower() != "nan" else "pcs",
                "warranty_months": warranty_months,
                "location_rack": location_rack,
                "position_bin": position_bin
            })
        except Exception as err:
            errors.append(f"Row {row_num} error: {err}")

    return items, errors
